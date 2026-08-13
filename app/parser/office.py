"""Office 文档解析:docx / pptx / xlsx / doc → Unit 列表。"""
import subprocess
import tempfile
from pathlib import Path

from app.models import Unit
from app.parser import _MAX_EMBEDDED_IMAGES
from app.parser.images import describe_image, ocr_image


def parse_doc(path: str | Path) -> list[Unit]:
    """老式 .doc 二进制 Word:经 LibreOffice 无头转换提取文本。

    依赖系统 soffice 命令(未安装时抛 DOC_CONVERT_FAILED,由上层容错跳过)。
    """
    import shutil

    if shutil.which("soffice") is None:
        raise ValueError("DOC_CONVERT_FAILED: soffice 未安装")
    with tempfile.TemporaryDirectory(prefix="doc2txt-") as tmp_dir:
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "txt:Text (encoded):UTF8",
             "--outdir", tmp_dir, str(path)],
            capture_output=True, timeout=180,
        )
        out = Path(tmp_dir) / (Path(path).stem + ".txt")
        if result.returncode != 0 or not out.exists():
            raise ValueError("DOC_CONVERT_FAILED")
        text = out.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    return [Unit(material_id=0, kind="text", content=text)]


def _parse_image_units(image_bytes: bytes) -> Unit:
    try:
        ocr_text = ocr_image(image_bytes)
        description = describe_image(image_bytes, ocr_text)
    except Exception:
        ocr_text, description = "", ""
    return Unit(material_id=0, kind="image", content=ocr_text, image_desc=description)


def parse_docx(path: str | Path) -> list[Unit]:
    from docx import Document

    doc = Document(path)
    units: list[Unit] = []
    paragraph_index = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        paragraph_index += 1
        style = (para.style.name or "").lower()
        kind = "heading" if style.startswith("heading") or style.startswith("标题") else "text"
        units.append(Unit(material_id=0, kind=kind, content=text, paragraph=paragraph_index))
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        content = "\n".join(" | ".join(row) for row in rows)
        units.append(Unit(material_id=0, kind="table", content=content, paragraph=paragraph_index + 1))
    image_count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            image_count += 1
            if image_count > _MAX_EMBEDDED_IMAGES:
                continue  # 超过上限的图片跳过,不阻塞主流程
            units.append(_parse_image_units(rel.target_part.blob))
    return units


def parse_pptx(path: str | Path) -> list[Unit]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    units: list[Unit] = []
    image_count = 0
    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(
                    p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
                )
                if text:
                    units.append(Unit(material_id=0, kind="text", content=text, page=slide_index))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
                if image_count > _MAX_EMBEDDED_IMAGES:
                    continue
                image_unit = _parse_image_units(shape.image.blob)
                image_unit.page = slide_index
                units.append(image_unit)
    return units


def parse_xlsx(path: str | Path) -> list[Unit]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    units: list[Unit] = []
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                values = ["" if v is None else str(v) for v in row]
                if any(values):
                    rows.append(" | ".join(values))
            if rows:
                content = f"[Sheet: {ws.title}]\n" + "\n".join(rows)
                units.append(Unit(material_id=0, kind="table", content=content))
    finally:
        wb.close()
    return units
