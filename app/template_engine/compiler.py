"""Deterministic DOCX template compiler.

The compiler intentionally separates two tracks:
- JSON schema: semantic roles, computed style, placeholders, and QA.
- Original DOCX: kept as the rendering base so complex OOXML is not re-created.
"""
from __future__ import annotations

import hashlib
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Length
from docx.oxml.ns import qn

COMPILER_VERSION = "template-compiler-2.0"

ROLE_DOCUMENT_TITLE = "document_title"
ROLE_HEADING_1 = "heading_1"
ROLE_HEADING_2 = "heading_2"
ROLE_HEADING_3 = "heading_3"
ROLE_BODY = "body"
ROLE_METADATA_LABEL = "metadata_label"
ROLE_METADATA_VALUE = "metadata_value"
ROLE_SIGNATURE = "signature"

_CJK_HEADING_1 = re.compile(r"^(?:[一二三四五六七八九十]+[、.]|第[一二三四五六七八九十\d]+[章节部分])\s*\S+")
_CJK_HEADING_2 = re.compile(r"^[（(][一二三四五六七八九十\d]+[)）]\s*\S+")
_NUM_HEADING_1 = re.compile(r"^\d+(?:\.(?!\d)|[、\s])\s*\S+")
_NUM_HEADING_2 = re.compile(r"^\d+\.\d+(?:\.(?!\d)|[、\s])\s*\S+")
_NUM_HEADING_3 = re.compile(r"^\d+\.\d+\.\d+[、\s]?\s*\S+")
_NUM_HEADING = re.compile(r"^\d+(?:\.\d+)*[.、]?\s*\S+")
_PLACEHOLDER = re.compile(r"(\{\{[^}]+\}\}|《[^》]+》|【[^】]+】|________+|_{4,})")


def compile_template(path) -> dict:
    """Compile a DOCX into a structured template schema.

    No LLM is used here. Unknown or unobservable details remain explicit instead
    of being guessed, which lets the UI and conformance checker show gaps.
    """
    source = Path(path)
    if not source.exists() or source.suffix.lower() != ".docx":
        return {}
    doc = Document(source)
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    role_paragraphs = _detect_role_paragraphs(paragraphs)
    schema = {
        "template_version": "1.0",
        "compiler_version": COMPILER_VERSION,
        "compiled_at": datetime.now(timezone.utc).isoformat(),
        "source": {
            "path": str(source),
            "filename": source.name,
            "sha256": _sha256(source),
        },
        "document": _document_layout(doc),
        "structure": _document_structure(doc, role_paragraphs),
        "style": {
            "roles": {
                role: _computed_role_style(doc, role, para)
                for role, para in role_paragraphs.items()
                if para is not None
            },
            "numbering": _numbering(doc),
            "tables": _tables(doc),
            "headers": _headers(doc),
            "footers": _footers(doc),
            "objects": _objects(doc),
        },
        "components": _components(doc, role_paragraphs),
    }
    schema["quality"] = _quality(schema)
    return schema


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _document_layout(doc) -> dict:
    sections = []
    for section in doc.sections:
        sections.append({
            "page": {
                "width_cm": _cm(section.page_width),
                "height_cm": _cm(section.page_height),
                "orientation": "landscape" if section.orientation == WD_ORIENT.LANDSCAPE else "portrait",
            },
            "margins": {
                "top_cm": _cm(section.top_margin),
                "bottom_cm": _cm(section.bottom_margin),
                "left_cm": _cm(section.left_margin),
                "right_cm": _cm(section.right_margin),
                "gutter_cm": _cm(getattr(section, "gutter", None)),
                "header_distance_cm": _cm(section.header_distance),
                "footer_distance_cm": _cm(section.footer_distance),
            },
            "header_footer": {
                "different_first_page": bool(section.different_first_page_header_footer),
                "header_text": _header_footer_text(section.header),
                "footer_text": _header_footer_text(section.footer),
            },
        })
    first = sections[0] if sections else {}
    return {
        "page": first.get("page", {}),
        "margins": first.get("margins", {}),
        "sections": sections,
        "defaults": _doc_defaults(doc),
    }


def _document_structure(doc, role_paragraphs: dict) -> dict:
    headings = []
    for paragraph_index, para in enumerate(doc.paragraphs):
        if para._p is getattr(role_paragraphs.get(ROLE_DOCUMENT_TITLE), "_p", None):
            continue
        for item in _heading_items(para):
            item["paragraph_index"] = paragraph_index
            headings.append(item)
    body_anchor = next((item for item in headings if item.get("paragraph_index") is not None), None)
    body_index = (
        body_anchor.get("paragraph_index")
        if body_anchor
        else _paragraph_index(doc, role_paragraphs.get(ROLE_BODY))
    )
    title_para = role_paragraphs.get(ROLE_DOCUMENT_TITLE)
    return {
        "roles_detected": sorted(k for k, v in role_paragraphs.items() if v is not None),
        "heading_patterns": headings[:20],
        "heading_tree": _heading_tree(headings),
        "template_instructions": _template_instructions(doc),
        "metadata_fields": _metadata_fields(doc),
        "signature_fields": _signature_fields(doc),
        "placeholders": _placeholder_registry(doc, role_paragraphs),
        "render_contract": {
            "mode": "replace_title_insert_body",
            "title_anchor": {
                "location": "paragraph",
                "paragraph_index": _paragraph_index(doc, title_para),
            } if title_para is not None else None,
            "body_anchor": {
                "location": "paragraph",
                "paragraph_index": body_index,
            } if body_index is not None else None,
            "preserve_prefix": bool(title_para is not None and body_index is not None),
        },
    }


def _detect_role_paragraphs(paragraphs: list) -> dict:
    roles = {
        ROLE_DOCUMENT_TITLE: None,
        ROLE_HEADING_1: None,
        ROLE_HEADING_2: None,
        ROLE_HEADING_3: None,
        ROLE_BODY: None,
        ROLE_SIGNATURE: None,
    }
    title = _best_title_paragraph(paragraphs[:8])
    roles[ROLE_DOCUMENT_TITLE] = title
    body_candidates = [para for para in paragraphs if para is not title]
    roles[ROLE_BODY] = _best_body_paragraph(body_candidates)
    for para in body_candidates:
        role = _paragraph_role(para)
        if role in {ROLE_HEADING_1, ROLE_HEADING_2, ROLE_HEADING_3} and roles[role] is None:
            roles[role] = para
    for para in body_candidates:
        text = para.text.strip()
        if roles[ROLE_SIGNATURE] is None and re.search(r"(单位|日期|年\s*月\s*日|盖章|署名)", text):
            roles[ROLE_SIGNATURE] = para
    if roles[ROLE_BODY] is None and body_candidates:
        roles[ROLE_BODY] = next((p for p in body_candidates if _paragraph_role(p) == ROLE_BODY), body_candidates[0])
    return roles


def _best_title_paragraph(paragraphs: list):
    """Choose a title from the cover area without assuming the first text is it."""
    scored = []
    for index, para in enumerate(paragraphs):
        text = para.text.strip()
        if not text or _looks_like_non_body_meta(text) or _is_template_instruction(text):
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        align = _alignment(para.alignment, _pf_value(para.style.paragraph_format if para.style is not None else None, "alignment"))
        run = next((item for item in para.runs if item.text.strip()), None)
        size = getattr(getattr(run, "font", None), "size", None)
        score = max(0, 7 - index)
        if align == "center":
            score += 8
        if any(marker in style for marker in ("title", "标题", "题名")):
            score += 8
        if size is not None and size.pt >= 16:
            score += 5
        if 4 <= len(text) <= 60:
            score += 3
        if _paragraph_role(para) in {ROLE_HEADING_2, ROLE_HEADING_3}:
            score -= 10
        scored.append((score, -index, para))
    return max(scored, key=lambda item: (item[0], item[1]))[2] if scored else (paragraphs[0] if paragraphs else None)


def _best_body_paragraph(paragraphs: list):
    candidates = [
        (_body_candidate_score(para), para)
        for para in paragraphs
        if _paragraph_role(para) == ROLE_BODY and _looks_like_body_sample(para)
    ]
    candidates = [(score, para) for score, para in candidates if score > 0]
    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0][1]


def _body_candidate_score(para) -> int:
    text = para.text.strip()
    style = para.style.name if para.style is not None else ""
    align = _alignment(
        para.alignment,
        _pf_value(para.style.paragraph_format if para.style is not None else None, "alignment"),
        None,
    )
    score = 0
    if "正文" in style or "body" in style.lower():
        score += 8
    if len(text) >= 24:
        score += 5
    if len(text) >= 48:
        score += 3
    if _is_template_instruction(text):
        score += 4
    if align in {None, "both", "justify"}:
        score += 5
    if align in {"center", "right"}:
        score -= 8
    if _looks_like_non_body_meta(text):
        score -= 20
    if re.search(r"(报告|方案|综述|总结)$", text) and len(text) <= 28:
        score -= 10
    return score


def _paragraph_role(para) -> str:
    text = para.text.strip()
    style = (para.style.name or "").lower() if para.style is not None else ""
    semantic_level = _semantic_heading_level(para)
    if semantic_level == 1 or _CJK_HEADING_1.match(text) or _NUM_HEADING_1.match(text) or style.startswith(("heading 1", "标题 1")) or "一级标题" in style:
        return ROLE_HEADING_1
    if semantic_level == 2 or _CJK_HEADING_2.match(text) or _NUM_HEADING_2.match(text) or style.startswith(("heading 2", "标题 2")) or "二级标题" in style:
        return ROLE_HEADING_2
    if semantic_level == 3 or _NUM_HEADING_3.match(text) or style.startswith(("heading 3", "标题 3")) or "三级标题" in style:
        return ROLE_HEADING_3
    return ROLE_BODY


def _computed_role_style(doc, role: str, para) -> dict:
    run = next((r for r in para.runs if r.text.strip()), para.runs[0] if para.runs else None)
    return {
        "paragraph": _paragraph_style(doc, para),
        "run": _run_style(doc, para, run),
        "sample_text": _role_sample_text(role, para.text),
        "samples": _role_samples(role, para.text),
        "source_style": para.style.name if para.style is not None else "",
        "source_index": _paragraph_index(doc, para),
        "confidence": _role_confidence(role, para),
        "computed_from": ["style_definition", "paragraph_direct_formatting", "run_direct_formatting"],
    }


def _paragraph_style(doc, para) -> dict:
    pf = para.paragraph_format
    style_pf = para.style.paragraph_format if para.style is not None else None
    normal_pf = doc.styles["Normal"].paragraph_format
    line_spacing = _line_spacing(pf.line_spacing, _pf_value(style_pf, "line_spacing"), normal_pf.line_spacing)
    return {
        "alignment": _alignment(para.alignment, _pf_value(style_pf, "alignment"), normal_pf.alignment),
        "left_indent_cm": _length_value(pf.left_indent, _pf_value(style_pf, "left_indent"), normal_pf.left_indent),
        "right_indent_cm": _length_value(pf.right_indent, _pf_value(style_pf, "right_indent"), normal_pf.right_indent),
        "first_line_indent_cm": _length_value(pf.first_line_indent, _pf_value(style_pf, "first_line_indent"), normal_pf.first_line_indent),
        "space_before_pt": _default_number(_pt_value(pf.space_before, _pf_value(style_pf, "space_before"), normal_pf.space_before), 0.0),
        "space_after_pt": _default_number(_pt_value(pf.space_after, _pf_value(style_pf, "space_after"), normal_pf.space_after), 0.0),
        "line_spacing": {"type": "multiple", "value": 1.0} if line_spacing == "unknown" else line_spacing,
        "keep_with_next": _bool_value(pf.keep_with_next, _pf_value(style_pf, "keep_with_next"), normal_pf.keep_with_next),
        "keep_together": _bool_value(pf.keep_together, _pf_value(style_pf, "keep_together"), normal_pf.keep_together),
        "page_break_before": _bool_value(pf.page_break_before, _pf_value(style_pf, "page_break_before"), normal_pf.page_break_before),
        "widow_control": _widow_control(para),
        "contextual_spacing": _contextual_spacing(para),
        "tabs": _tabs(pf),
        "outline_level": _outline_level(para),
        "shading": _shading(para._p),
        "borders": _computed_borders(para),
    }


def _run_style(doc, para, run) -> dict:
    style_font = para.style.font if para.style is not None else None
    normal_font = doc.styles["Normal"].font
    font = run.font if run is not None else None
    xmlfont = _rpr_font_values(para.style) or {}
    defaults = _doc_default_rpr_values(doc)
    for key, value in defaults.items():
        if not xmlfont.get(key):
            xmlfont[key] = value
    return {
        "font_east_asia": _font_name(font, style_font, normal_font, "eastAsia", xmlfont),
        "font_ascii": _font_name(font, style_font, normal_font, "ascii", xmlfont),
        "font_hansi": _font_name(font, style_font, normal_font, "hAnsi", xmlfont),
        "font_cs": _font_name(font, style_font, normal_font, "cs", xmlfont),
        "font_size_pt": _font_size(font, style_font, normal_font, xmlfont),
        "bold": _bool_value(getattr(font, "bold", None), getattr(style_font, "bold", None), getattr(normal_font, "bold", None)),
        "italic": _bool_value(getattr(font, "italic", None), getattr(style_font, "italic", None), getattr(normal_font, "italic", None)),
        "underline": _bool_value(getattr(font, "underline", None), getattr(style_font, "underline", None), getattr(normal_font, "underline", None)),
        "strike": _bool_value(getattr(font, "strike", None), getattr(style_font, "strike", None), getattr(normal_font, "strike", None)),
        "font_color": _font_color(font, style_font, normal_font),
        "highlight": str(getattr(font, "highlight_color", "") or "") if font is not None else "unknown",
        "character_spacing": _run_prop(run, "spacing"),
        "scaling": _run_prop(run, "w"),
        "position": _run_prop(run, "position"),
        "superscript": _bool_value(getattr(font, "superscript", None), None, None),
        "subscript": _bool_value(getattr(font, "subscript", None), None, None),
        "small_caps": _bool_value(getattr(font, "small_caps", None), None, None),
        "all_caps": _bool_value(getattr(font, "all_caps", None), None, None),
        "language": _run_language(run),
    }


def _doc_defaults(doc) -> dict:
    normal = doc.styles["Normal"]
    xmlfont = _rpr_font_values(normal) or {}
    for key, value in _doc_default_rpr_values(doc).items():
        if not xmlfont.get(key):
            xmlfont[key] = value
    return {
        "run": {
            "font_east_asia": _font_name(None, normal.font, None, "eastAsia", xmlfont),
            "font_ascii": _font_name(None, normal.font, None, "ascii", xmlfont),
            "font_size_pt": _font_size(None, normal.font, None, xmlfont),
        },
        "paragraph": {
            "line_spacing": _line_spacing(normal.paragraph_format.line_spacing),
            "first_line_indent_cm": _length_value(normal.paragraph_format.first_line_indent),
        },
    }


def _numbering(doc) -> dict:
    patterns = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        pattern = ""
        if _CJK_HEADING_1.match(text):
            pattern = "cjk_level_1"
        elif _CJK_HEADING_2.match(text):
            pattern = "cjk_level_2"
        elif _NUM_HEADING_3.match(text):
            pattern = "arabic_level_3"
        elif _NUM_HEADING_2.match(text):
            pattern = "arabic_level_2"
        elif _NUM_HEADING_1.match(text):
            pattern = "arabic_level_1"
        if pattern and pattern not in [p["format"] for p in patterns]:
            sample = next((line for line in _logical_lines(text) if _heading_items(line)), text)
            patterns.append({"format": pattern, "sample": _text_pattern(sample)})
    definitions = _numbering_definitions(doc)
    return {
        "patterns": patterns,
        "definitions": definitions,
        "role_levels": _numbering_role_levels(doc, definitions),
    }


def _numbering_definitions(doc) -> list[dict]:
    """Decode OOXML numbering instead of relying on rendered paragraph text."""
    try:
        root = doc.part.numbering_part.element
    except Exception:
        return []
    abstract_by_id = {
        node.get(qn("w:abstractNumId")): node
        for node in root.findall(qn("w:abstractNum"))
    }
    definitions = []
    for num in root.findall(qn("w:num")):
        num_id = num.get(qn("w:numId"))
        abstract_ref = num.find(qn("w:abstractNumId"))
        abstract_id = abstract_ref.get(qn("w:val")) if abstract_ref is not None else None
        abstract = abstract_by_id.get(abstract_id)
        if abstract is None:
            continue
        levels = []
        for level in abstract.findall(qn("w:lvl")):
            levels.append(_numbering_level(level))
        if levels:
            definitions.append({
                "num_id": _as_int(num_id),
                "abstract_num_id": _as_int(abstract_id),
                "levels": levels,
            })
    return definitions


def _numbering_level(level) -> dict:
    def value(name: str, default=""):
        node = level.find(qn(f"w:{name}"))
        return node.get(qn("w:val"), default) if node is not None else default

    ppr = level.find(qn("w:pPr"))
    ind = ppr.find(qn("w:ind")) if ppr is not None else None
    return {
        "level": _as_int(level.get(qn("w:ilvl"))),
        "start": _as_int(value("start", "1")),
        "number_format": value("numFmt"),
        "level_text": value("lvlText"),
        "suffix": value("suff", "tab"),
        "paragraph_style": value("pStyle"),
        "alignment": value("lvlJc"),
        "left_twips": _as_int(ind.get(qn("w:left"))) if ind is not None else None,
        "hanging_twips": _as_int(ind.get(qn("w:hanging"))) if ind is not None else None,
    }


def _numbering_role_levels(doc, definitions: list[dict]) -> dict:
    role_levels: dict[str, dict] = {}
    definitions_by_id = {item.get("num_id"): item for item in definitions}
    for para in doc.paragraphs:
        semantic = _semantic_heading_level(para)
        numbering = _paragraph_numbering(para)
        if semantic not in {1, 2, 3} or not numbering:
            continue
        definition = definitions_by_id.get(numbering.get("num_id")) or {}
        level = next(
            (item for item in definition.get("levels") or [] if item.get("level") == numbering.get("level")),
            {},
        )
        if level:
            role_levels[f"heading_{semantic}"] = {
                "num_id": numbering.get("num_id"),
                **level,
                "style_name": para.style.name if para.style is not None else "",
            }
    used_style_ids = {
        para.style.style_id
        for para in doc.paragraphs
        if para.text.strip() and para.style is not None and _semantic_heading_level(para) in {1, 2, 3}
    }
    for definition in definitions:
        for level in definition.get("levels") or []:
            style_id = str(level.get("paragraph_style") or "")
            if not style_id or style_id not in used_style_ids:
                continue
            style = next((item for item in doc.styles if item.style_id == style_id), None)
            style_name = style.name if style is not None else style_id
            semantic = _style_heading_level(style_name)
            if semantic and f"heading_{semantic}" not in role_levels:
                role_levels[f"heading_{semantic}"] = {
                    "num_id": definition.get("num_id"),
                    **level,
                    "style_name": style_name,
                }
    return role_levels


def _paragraph_numbering(para) -> dict:
    try:
        style = para.style
        ppr = para._p.pPr
        num_pr = ppr.numPr if ppr is not None else None
        visited = set()
        while num_pr is None and style is not None and style.style_id not in visited:
            visited.add(style.style_id)
            ppr = style.element.pPr
            num_pr = ppr.numPr if ppr is not None else None
            style = style.base_style
        if num_pr is None:
            return {}
        return {
            "num_id": _as_int(num_pr.numId.val) if num_pr.numId is not None else None,
            "level": _as_int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0,
        }
    except Exception:
        return {}


def _semantic_heading_level(para) -> int:
    style_name = para.style.name if para.style is not None else ""
    named_level = _style_heading_level(style_name)
    if named_level:
        return named_level
    outline = _outline_level(para)
    if str(outline).isdigit() and 0 <= int(outline) <= 2:
        return int(outline) + 1
    numbering_level = _style_numbering_level(para)
    return numbering_level if numbering_level in {1, 2, 3} else 0


def _style_heading_level(style_name: str) -> int:
    value = str(style_name or "").lower().replace("_", " ")
    patterns = {
        1: ("heading 1", "标题 1", "标题1", "一级标题", "标题一"),
        2: ("heading 2", "标题 2", "标题2", "二级标题", "标题二"),
        3: ("heading 3", "标题 3", "标题3", "三级标题", "标题三"),
    }
    for level, markers in patterns.items():
        if any(marker in value for marker in markers):
            return level
    return 0


def _style_numbering_level(para) -> int:
    style_id = getattr(getattr(para, "style", None), "style_id", "")
    if not style_id:
        return 0
    try:
        root = para.part.numbering_part.element
        for abstract in root.findall(qn("w:abstractNum")):
            for level in abstract.findall(qn("w:lvl")):
                pstyle = level.find(qn("w:pStyle"))
                if pstyle is not None and pstyle.get(qn("w:val")) == style_id:
                    return (_as_int(level.get(qn("w:ilvl"))) or 0) + 1
    except Exception:
        pass
    return 0


def _as_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _default_number(value, fallback: float):
    return fallback if value in (None, "", "unknown") else value


def _tables(doc) -> list[dict]:
    tables = []
    for idx, table in enumerate(doc.tables[:10]):
        tables.append({
            "index": idx,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "style_name": table.style.name if table.style is not None else "",
            "column_widths_cm": [_cm(cell.width) for cell in table.rows[0].cells] if table.rows else [],
            "cells": _table_cells(table),
        })
    return tables


def _table_cells(table) -> list[dict]:
    cells = []
    for r_idx, row in enumerate(table.rows[:5]):
        for c_idx, cell in enumerate(row.cells[:8]):
            text = "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            cells.append({
                "row": r_idx,
                "col": c_idx,
                "text_pattern": _text_pattern(text),
                "vertical_alignment": str(cell.vertical_alignment) if cell.vertical_alignment else "",
                "paragraph": _paragraph_style_from_any(cell.paragraphs[0]) if cell.paragraphs else {},
                "borders": _tc_borders(cell),
                "shading": _shading(cell._tc),
            })
    return cells


def _headers(doc) -> list[dict]:
    return [{"section": idx, "text": _header_footer_text(section.header), "paragraph_count": len(section.header.paragraphs)}
            for idx, section in enumerate(doc.sections)]


def _footers(doc) -> list[dict]:
    return [{"section": idx, "text": _header_footer_text(section.footer), "paragraph_count": len(section.footer.paragraphs)}
            for idx, section in enumerate(doc.sections)]


def _objects(doc) -> dict:
    return {
        "inline_images": len(getattr(doc, "inline_shapes", [])),
        "hyperlinks": _count_xpath(doc.element, ".//w:hyperlink"),
        "bookmarks": _count_xpath(doc.element, ".//w:bookmarkStart"),
        "fields": _count_xpath(doc.element, ".//w:fldChar"),
        "textboxes": _count_xpath(doc.element, ".//w:txbxContent"),
    }


def _components(doc, role_paragraphs: dict) -> list[dict]:
    components = []
    if role_paragraphs.get(ROLE_DOCUMENT_TITLE) is not None:
        components.append({"type": "TITLE_BLOCK", "role": ROLE_DOCUMENT_TITLE, "placeholder": "report_title"})
    metadata = _metadata_fields(doc)
    if metadata:
        components.append({"type": "METADATA_BLOCK", "fields": metadata})
    if doc.tables:
        components.append({"type": "TABLE", "count": len(doc.tables)})
    if any(_header_footer_text(s.header) for s in doc.sections):
        components.append({"type": "HEADER"})
    if any(_header_footer_text(s.footer) for s in doc.sections):
        components.append({"type": "FOOTER"})
    if role_paragraphs.get(ROLE_SIGNATURE) is not None:
        components.append({"type": "SIGNATURE_BLOCK", "role": ROLE_SIGNATURE, "fields": _signature_fields(doc)})
    components.append({"type": "BODY_SECTION", "role": ROLE_BODY, "placeholder": "report_body"})
    return components


def _quality(schema: dict) -> dict:
    roles = schema.get("style", {}).get("roles", {})
    required = {
        "主标题": ROLE_DOCUMENT_TITLE,
        "正文": ROLE_BODY,
        "一级标题": ROLE_HEADING_1,
    }
    checks = []
    for label, role in required.items():
        role_style = roles.get(role, {})
        checks.append({"item": f"识别{label}", "status": "ok" if role_style else "unknown"})
        checks.append({"item": f"{label}字体", "status": "ok" if role_style.get("run", {}).get("font_east_asia") not in ("", "unknown", None) else "unknown"})
        checks.append({"item": f"{label}字号", "status": "ok" if role_style.get("run", {}).get("font_size_pt") not in ("", "unknown", None) else "unknown"})
    body_para = roles.get(ROLE_BODY, {}).get("paragraph", {})
    for field in ("first_line_indent_cm", "line_spacing", "space_after_pt"):
        checks.append({"item": f"正文{field}", "status": "ok" if body_para.get(field) not in ("unknown", None, "") else "unknown"})
    numbering = schema.get("style", {}).get("numbering", {})
    checks.append({
        "item": "标题编号",
        "status": "ok" if numbering.get("patterns") or numbering.get("definitions") else "not_applicable",
    })
    contract = schema.get("structure", {}).get("render_contract", {})
    checks.append({
        "item": "原模板填充锚点",
        "status": "ok" if contract.get("title_anchor") and contract.get("body_anchor") else "unknown",
    })
    unknown = [c["item"] for c in checks if c["status"] == "unknown"]
    applicable = [item for item in checks if item["status"] != "not_applicable"]
    return {
        "checks": checks,
        "unknown_items": unknown,
        "completeness": round((len(applicable) - len(unknown)) / max(len(applicable), 1), 2),
        "render_ready": not any(
            item["status"] == "unknown" and item["item"] in {"识别主标题", "识别正文", "识别一级标题", "原模板填充锚点"}
            for item in checks
        ),
        "recommendation": "可用于正式导出" if not unknown else "部分模板属性需人工复核",
    }


def _heading_items(para_or_text) -> list[dict]:
    """识别标题项:优先用 Word 样式(中英文兼容),否则用增强正则兜底。
    兼容传 para 或纯文本两种调用。"""
    items = []
    if hasattr(para_or_text, "text"):
        para = para_or_text
        text = para.text or ""
        style_name = ((para.style.name or "").lower() if para.style is not None else "")
    else:
        para = None
        text = para_or_text or ""
        style_name = ""
    role = ""
    if para is not None:
        level = _semantic_heading_level(para)
        role = {1: ROLE_HEADING_1, 2: ROLE_HEADING_2, 3: ROLE_HEADING_3}.get(level, "")
    for line in _logical_lines(text):
        r = role
        if not r:
            if _CJK_HEADING_1.match(line) or _NUM_HEADING_1.match(line):
                r = ROLE_HEADING_1
            elif _CJK_HEADING_2.match(line) or _NUM_HEADING_2.match(line):
                r = ROLE_HEADING_2
            elif _NUM_HEADING_3.match(line):
                r = ROLE_HEADING_3
        if r:
            item = {"role": r, "level": _heading_level(r), "text_pattern": _text_pattern(line)}
            if para is not None:
                numbering = _paragraph_numbering(para)
                if numbering:
                    item["numbering"] = numbering
            items.append(item)
    return items


def _heading_level(role: str) -> int:
    return {ROLE_HEADING_1: 1, ROLE_HEADING_2: 2, ROLE_HEADING_3: 3}.get(role, 9)


def _heading_tree(headings: list[dict]) -> list[dict]:
    """Build a heading hierarchy for UI and export policy inspection."""
    roots: list[dict] = []
    stack: list[dict] = []
    for item in headings:
        level = int(item.get("level") or _heading_level(item.get("role", "")))
        node = {
            "role": item.get("role", ""),
            "level": level,
            "text_pattern": item.get("text_pattern", ""),
            "children": [],
        }
        while stack and int(stack[-1].get("level") or 9) >= level:
            stack.pop()
        if stack:
            stack[-1].setdefault("children", []).append(node)
        else:
            roots.append(node)
        stack.append(node)
    return roots[:30]


def _looks_like_body_sample(para) -> bool:
    text = para.text.strip()
    style = para.style.name if para.style is not None else ""
    if not text or _heading_items(text):
        return False
    if _looks_like_non_body_meta(text):
        return False
    if _is_template_instruction(text):
        return "正文样式" in style
    if any(marker in text for marker in ("摘", "关键词", "中图法分类号")) and ("[" in text or "：" in text):
        return False
    if len(text) < 12 and "正文" not in style:
        return False
    return True


def _looks_like_non_body_meta(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return True
    if "××" in stripped or "XXX" in stripped.upper():
        return True
    if re.fullmatch(r"[（(]?\d{4}\s*年度[）)]?", stripped):
        return True
    if re.fullmatch(r"[×xX]{1,4}\s*年\s*[×xX]?\s*月\s*[×xX]?\s*日", stripped):
        return True
    if re.fullmatch(r".{1,20}(单位|部门|委员会|办公室)", stripped):
        return True
    if re.fullmatch(r"[（(].{1,30}[）)]", stripped):
        return True
    return False


def _is_template_instruction(text: str) -> bool:
    stripped = text.strip()
    return bool(
        (stripped.startswith("[") and stripped.endswith("]"))
        or re.search(r"[【\[]\s*(?:填写|说明|关键词|占位|按需)[^】\]]*[】\]]", stripped)
        or ("填写" in stripped and "按《" in stripped)
        or ("填写" in stripped and ("[" in stripped or "。" in stripped))
        or stripped.startswith(("示例：", "示例:", "范例：", "范例:"))
        or "说明" in stripped
    )


def _template_instructions(doc) -> list[dict]:
    instructions = []
    for paragraph_index, para in enumerate(doc.paragraphs):
        chunks = _split_heading_instruction_chunks(para.text)
        for chunk in chunks:
            if chunk["role"] == "TEMPLATE_INSTRUCTION":
                instructions.append({
                    "role": "TEMPLATE_INSTRUCTION",
                    "text_pattern": _text_pattern(chunk["text"]),
                    "location": "paragraph",
                    "paragraph_index": paragraph_index,
                    "render": False,
                })
    return instructions[:50]


def _role_sample_text(role: str, text: str) -> str:
    samples = _role_samples(role, text)
    if samples:
        return samples[0][:120]
    return text.strip()[:120]


def _role_samples(role: str, text: str) -> list[str]:
    if role in {ROLE_HEADING_1, ROLE_HEADING_2, ROLE_HEADING_3}:
        return [item["text_pattern"] for item in _heading_items(text) if item["role"] == role] or [text.strip()[:120]]
    if role == ROLE_BODY:
        return [
            chunk["text"][:120]
            for chunk in _split_heading_instruction_chunks(text)
            if chunk["role"] == "TEMPLATE_INSTRUCTION" or not _heading_items(chunk["text"])
        ][:3] or [text.strip()[:120]]
    return [text.strip()[:120]]


def _split_heading_instruction_chunks(text: str) -> list[dict]:
    lines = _logical_lines(text)
    chunks = []
    current_heading = None
    for line in lines:
        heading = _heading_items(line)
        if heading:
            current_heading = heading[0]["role"]
            chunks.append({"role": current_heading, "text": line})
        else:
            chunks.append({
                "role": "TEMPLATE_INSTRUCTION" if current_heading or _is_template_instruction(line) else ROLE_BODY,
                "text": line,
            })
    return chunks


def _logical_lines(text: str) -> list[str]:
    lines = []
    for raw in re.split(r"[\r\n]+", text or ""):
        line = raw.strip()
        if line:
            lines.append(line)
    return lines


def _metadata_fields(doc) -> list[dict]:
    fields = []
    labels = ("报告单位", "报告时间", "报告主题", "报告编号")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                for label in labels:
                    if label in text:
                        fields.append({
                            "label": label,
                            "placeholder": _placeholder_name(label, "metadata"),
                            "component": "METADATA_BLOCK",
                            "text_pattern": _text_pattern(text),
                        })
                        break
    deduped = []
    seen = set()
    for field in fields:
        key = (field["label"], field["text_pattern"])
        if key not in seen:
            seen.add(key)
            deduped.append(field)
    return deduped[:20]


def _signature_fields(doc) -> list[dict]:
    fields = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if re.search(r"(×××单位|盖章|单位)", text) and re.search(r"(×年×月×日|年\s*月\s*日|日期)", text):
            fields.append({
                "placeholder": "signature_unit",
                "component": "SIGNATURE_BLOCK",
                "text_pattern": _text_pattern(text),
            })
            fields.append({
                "placeholder": "signature_date",
                "component": "SIGNATURE_BLOCK",
                "text_pattern": _text_pattern(text),
            })
        elif re.search(r"(×××单位|盖章|单位)", text):
            fields.append({"placeholder": "signature_unit", "component": "SIGNATURE_BLOCK", "text_pattern": _text_pattern(text)})
        elif re.search(r"(×年×月×日|年\s*月\s*日|日期)", text):
            fields.append({"placeholder": "signature_date", "component": "SIGNATURE_BLOCK", "text_pattern": _text_pattern(text)})
    deduped = []
    seen = set()
    for field in fields:
        key = (field["placeholder"], field["text_pattern"])
        if key not in seen:
            seen.add(key)
            deduped.append(field)
    return deduped[:8]


def _physical_placeholders(doc) -> list[dict]:
    found = []
    for para in doc.paragraphs:
        for match in _PLACEHOLDER.findall(para.text):
            found.append({"text": match, "location": "paragraph"})
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                for match in _PLACEHOLDER.findall(cell.text):
                    found.append({"text": match, "location": "table", "table": t_idx, "row": r_idx, "col": c_idx})
    return found[:50]


def _placeholder_registry(doc, role_paragraphs: dict) -> dict:
    registry = {}
    _register_placeholder(registry, "report_title", "text", "TITLE_BLOCK", True)
    _register_placeholder(registry, "report_body", "dynamic_content", "BODY_SECTION", True)
    for field in _metadata_fields(doc):
        _register_placeholder(registry, field["placeholder"], "text", "METADATA_BLOCK", False, label=field.get("label", ""))
    for field in _signature_fields(doc):
        _register_placeholder(registry, field["placeholder"], "text", "SIGNATURE_BLOCK", False)
    physical = _physical_placeholders(doc)
    for item in physical:
        key = _normalize_placeholder(item["text"])
        if not key:
            continue
        _register_placeholder(registry, key, "text", "UNKNOWN", False)
        registry[key]["physical_placeholder"] = item
    return registry


def _register_placeholder(registry: dict, key: str, typ: str, component: str, required: bool, label: str = "") -> None:
    if key not in registry:
        registry[key] = {
            "type": typ,
            "component": component,
            "required": required,
            "bindings": [{"component": component, "required": required}],
        }
        if label:
            registry[key]["label"] = label
        return
    item = registry[key]
    item["required"] = bool(item.get("required")) or required
    if item.get("component") in ("UNKNOWN", "METADATA_BLOCK") and component in ("TITLE_BLOCK", "BODY_SECTION"):
        item["component"] = component
    bindings = item.setdefault("bindings", [])
    if not any(b.get("component") == component for b in bindings):
        bindings.append({"component": component, "required": required})
    if label and not item.get("label"):
        item["label"] = label


def _placeholder_name(label: str, component: str = "metadata") -> str:
    mapping = {"报告单位": "report_unit", "报告时间": "report_date", "报告主题": "report_title", "报告编号": "report_no", "单位": "report_unit", "时间": "report_date", "日期": "report_date"}
    if component == "signature":
        mapping = {"单位": "signature_unit", "时间": "signature_date", "日期": "signature_date"}
    return mapping.get(label, label)


def _normalize_placeholder(text: str) -> str:
    value = text.strip("{}《》【】_ ").strip()
    mapping = {
        "报告标题": "report_title", "标题": "report_title", "报告主题": "report_title",
        "报告单位": "report_unit", "报告时间": "report_date", "报告日期": "report_date",
        "报告编号": "report_no", "正文": "report_body", "报告正文": "report_body",
    }
    if value in mapping:
        return mapping[value]
    if re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]{1,60}", value):
        return value
    return ""


def _header_footer_text(part) -> str:
    return "\n".join(p.text.strip() for p in part.paragraphs if p.text.strip())[:500]


def _paragraph_style_from_any(para) -> dict:
    pf = para.paragraph_format
    return {
        "alignment": _alignment(para.alignment, pf.alignment),
        "first_line_indent_cm": _length_value(pf.first_line_indent),
        "line_spacing": _line_spacing(pf.line_spacing),
        "space_after_pt": _pt_value(pf.space_after),
    }


def _paragraph_index(doc, para) -> int | None:
    if para is None:
        return None
    target = para._p
    return next((index for index, item in enumerate(doc.paragraphs) if item._p is target), None)


def _role_confidence(role: str, para) -> str:
    if para is None:
        return "unavailable"
    style_name = para.style.name if para.style is not None else ""
    if role == ROLE_DOCUMENT_TITLE:
        is_title_style = "title" in style_name.lower() or style_name in {"标题", "文档标题", "主标题"}
        return "high" if is_title_style else "medium"
    if role in {ROLE_HEADING_1, ROLE_HEADING_2, ROLE_HEADING_3}:
        expected = _heading_level(role)
        return "high" if _semantic_heading_level(para) == expected else "medium"
    if role == ROLE_BODY:
        return "high" if any(marker in style_name.lower() for marker in ("正文", "body", "normal")) else "medium"
    return "medium"


def _computed_borders(para) -> dict:
    direct = _p_borders(para._p)
    if any(v is not None for v in direct.values()):
        return direct
    if para.style is not None:
        style_borders = _p_borders(para.style.element)
        if any(v is not None for v in style_borders.values()):
            return style_borders
    return {"top": None, "bottom": None, "left": None, "right": None}


def _p_borders(element) -> dict:
    result = {"top": None, "bottom": None, "left": None, "right": None}
    try:
        ppr = element.pPr if hasattr(element, "pPr") else element.find(qn("w:pPr"))
        pbdr = ppr.find(qn("w:pBdr")) if ppr is not None else None
        if pbdr is None:
            return result
        for name in result:
            node = pbdr.find(qn(f"w:{name}"))
            if node is not None:
                result[name] = _border_node(node)
    except Exception:
        return result
    return result


def _tc_borders(cell) -> dict:
    result = {"top": None, "bottom": None, "left": None, "right": None}
    try:
        tcpr = cell._tc.tcPr
        borders = tcpr.find(qn("w:tcBorders")) if tcpr is not None else None
        if borders is None:
            return result
        for name in result:
            node = borders.find(qn(f"w:{name}"))
            if node is not None:
                result[name] = _border_node(node)
    except Exception:
        return result
    return result


def _border_node(node) -> dict:
    return {
        "style": node.get(qn("w:val"), ""),
        "width": node.get(qn("w:sz"), ""),
        "color": node.get(qn("w:color"), ""),
        "space": node.get(qn("w:space"), ""),
    }


def _shading(element) -> dict | str:
    try:
        target = element.pPr if hasattr(element, "pPr") else getattr(element, "tcPr", None)
        shd = target.find(qn("w:shd")) if target is not None else None
        if shd is None:
            return "none"
        return {"fill": shd.get(qn("w:fill"), ""), "color": shd.get(qn("w:color"), ""), "val": shd.get(qn("w:val"), "")}
    except Exception:
        return "unknown"


def _outline_level(para) -> str:
    try:
        ppr = para._p.pPr
        node = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        if node is not None:
            return node.get(qn("w:val"), "")
        style = para.style
        visited = set()
        while style is not None and style.style_id not in visited:
            visited.add(style.style_id)
            ppr = style.element.pPr
            node = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
            if node is not None:
                return node.get(qn("w:val"), "")
            style = style.base_style
    except Exception:
        pass
    return ""


def _widow_control(para) -> str | bool:
    try:
        ppr = para._p.pPr
        node = ppr.find(qn("w:widowControl")) if ppr is not None else None
        if node is None:
            return "unknown"
        return node.get(qn("w:val"), "1") != "0"
    except Exception:
        return "unknown"


def _contextual_spacing(para) -> str | bool:
    try:
        ppr = para._p.pPr
        node = ppr.find(qn("w:contextualSpacing")) if ppr is not None else None
        return bool(node is not None)
    except Exception:
        return "unknown"


def _tabs(pf) -> list[dict]:
    try:
        return [{"position_cm": _cm(stop.position), "alignment": str(stop.alignment), "leader": str(stop.leader)}
                for stop in pf.tab_stops]
    except Exception:
        return []


def _run_prop(run, prop: str) -> str:
    try:
        rpr = run._r.rPr if run is not None else None
        node = rpr.find(qn(f"w:{prop}")) if rpr is not None else None
        return node.get(qn("w:val"), "") if node is not None else "unknown"
    except Exception:
        return "unknown"


def _run_language(run) -> str:
    try:
        rpr = run._r.rPr if run is not None else None
        node = rpr.find(qn("w:lang")) if rpr is not None else None
        if node is None:
            return "unknown"
        return node.get(qn("w:eastAsia")) or node.get(qn("w:val")) or "unknown"
    except Exception:
        return "unknown"


def _font_name(font, style_font, normal_font, key: str, xmlfont=None) -> str:
    for candidate in (font, style_font, normal_font):
        if candidate is None:
            continue
        value = _rfont(candidate, key)
        if value:
            return value
        if key in {"ascii", "hAnsi"} and getattr(candidate, "name", None):
            return candidate.name
    # 兜底:python-docx 缓存读不到时,从样式 XML rPr 的 rFonts 解析(如 Normal 的宋体)
    if xmlfont:
        v = xmlfont.get(key) or xmlfont.get("eastAsia")
        if v:
            return v
    return "unknown"


def _rfont(font, key: str) -> str:
    try:
        element = font._element
        rfonts = element.rPr.rFonts if hasattr(element, "rPr") and element.rPr is not None else None
        if rfonts is None:
            return ""
        return rfonts.get(qn(f"w:{key}")) or ""
    except Exception:
        return ""


def _font_size(font, style_font, normal_font, xmlfont=None):
    for candidate in (font, style_font, normal_font):
        if candidate is None:
            continue
        size = getattr(candidate, "size", None)
        if size is not None:
            return round(size.pt, 1)
    if xmlfont and xmlfont.get("size_pt"):
        return xmlfont["size_pt"]
    return "unknown"


def _rpr_font_values(style):
    """从样式 XML rPr 解析实际字体(rFonts/sz)。python-docx 的 .font 缓存常读不到
    rPr 里的 rFonts/sz(如 Normal 样式定义的宋体/12pt),这里兜底,实现"忠实记录模板呈现"。"""
    if style is None:
        return None
    try:
        out = {"ascii": None, "hAnsi": None, "eastAsia": None, "cs": None, "size_pt": None}
        current = style
        visited = set()
        while current is not None and current.style_id not in visited:
            visited.add(current.style_id)
            rPr = current.element.find(qn("w:rPr"))
            rf = rPr.find(qn("w:rFonts")) if rPr is not None else None
            if rf is not None:
                for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                    if not out[key]:
                        out[key] = rf.get(qn("w:" + key))
            sz = rPr.find(qn("w:sz")) if rPr is not None else None
            if sz is not None and out["size_pt"] is None:
                try:
                    out["size_pt"] = round(int(sz.get(qn("w:val")) or 0) / 2, 1)
                except Exception:
                    pass
            current = current.base_style
        return out
    except Exception:
        return None


def _doc_default_rpr_values(doc) -> dict:
    out = {"ascii": None, "hAnsi": None, "eastAsia": None, "cs": None, "size_pt": None}
    try:
        defaults = doc.styles.element.find(qn("w:docDefaults"))
        rpr_default = defaults.find(qn("w:rPrDefault")) if defaults is not None else None
        rpr = rpr_default.find(qn("w:rPr")) if rpr_default is not None else None
        fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        if fonts is not None:
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                out[key] = fonts.get(qn("w:" + key))
        size = rpr.find(qn("w:sz")) if rpr is not None else None
        if size is not None:
            out["size_pt"] = round(int(size.get(qn("w:val")) or 0) / 2, 1)
    except Exception:
        pass
    return out


def _font_color(font, style_font, normal_font) -> str:
    for candidate in (font, style_font, normal_font):
        try:
            color = getattr(getattr(candidate, "color", None), "rgb", None)
            if color:
                return str(color)
        except Exception:
            continue
    return "000000"


def _alignment(*values) -> str:
    for value in values:
        if value is not None:
            text = str(value)
            if "CENTER" in text:
                return "center"
            if "RIGHT" in text:
                return "right"
            if "JUSTIFY" in text:
                return "justify"
            return "left"
    return "unknown"


def _length_value(*values):
    for value in values:
        if value is not None:
            return _cm(value)
    return "unknown"


def _pt_value(*values):
    for value in values:
        if value is not None:
            try:
                return round(value.pt, 1)
            except Exception:
                return "unknown"
    return "unknown"


def _line_spacing(*values):
    for value in values:
        if value is None:
            continue
        if isinstance(value, Length):
            return {"type": "exact_pt", "value": round(value.pt, 1)}
        if isinstance(value, (int, float)):
            return {"type": "multiple", "value": round(float(value), 2)}
        return {"type": "raw", "value": str(value)}
    return "unknown"


def _bool_value(*values):
    for value in values:
        if value is not None:
            return bool(value)
    return "unknown"


def _pf_value(pf, name: str):
    return getattr(pf, name, None) if pf is not None else None


def _cm(value):
    if value is None:
        return "unknown"
    try:
        return round(value.cm, 2)
    except Exception:
        return "unknown"


def _count_xpath(element, xpath: str) -> int:
    try:
        return len(element.xpath(xpath))
    except Exception:
        return 0


def _text_pattern(text: str) -> str:
    text = text.strip()
    if not text:
        return ""
    text = re.sub(r"\d{4}年\d{1,2}月\d{1,2}日", "{{date}}", text)
    text = re.sub(r"\d+", "{{number}}", text)
    text = _PLACEHOLDER.sub("{{placeholder}}", text)
    return text[:120]
