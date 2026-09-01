import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.export.docx import _prepare_render_body
from app.memory.style_profile import aggregate_style_metrics, select_exemplars
from app.rendering.headings import detect_numbering_strategy, format_heading
from app.template_engine.compiler import compile_template


class TemplateLearningTests(unittest.TestCase):
    def _save(self, doc: Document, name: str = "template.docx") -> Path:
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / name
        doc.save(path)
        return path

    def test_compiler_separates_heading_samples_from_instructions_and_body(self):
        doc = Document()
        title = doc.add_paragraph("年度工作总结报告")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.add_run().font.size = Pt(22)
        doc.add_paragraph("一、总体情况")
        h2 = doc.add_paragraph()
        h2.add_run("（一）重点工作推进情况\n说明主要工作内容。\n（二）重点成果形成情况\n总结阶段性成果。")
        doc.add_paragraph("围绕任务推进情况形成了可核验的工作记录和具体结果。")

        schema = compile_template(self._save(doc))
        roles = schema["style"]["roles"]

        self.assertEqual(roles["heading_2"]["samples"], ["（一）重点工作推进情况", "（二）重点成果形成情况"])
        self.assertNotIn("说明主要工作内容", roles["heading_2"]["sample_text"])
        self.assertIn("具体结果", roles["body"]["sample_text"])
        self.assertEqual(schema["style"]["numbering"]["patterns"][1]["sample"], "（一）重点工作推进情况")

    def test_compiler_does_not_treat_document_titles_as_placeholders(self):
        doc = Document()
        title = doc.add_paragraph("可信执行环境研究综述")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("中图法分类号按《中国图书馆分类法》填写")
        doc.add_paragraph("一、研究背景")
        doc.add_paragraph("本文围绕可信执行环境的远程证明机制开展系统分析。")

        placeholders = compile_template(self._save(doc))["structure"]["placeholders"]

        self.assertNotIn("中国图书馆分类法", placeholders)
        self.assertIn("report_title", placeholders)
        self.assertIn("report_body", placeholders)

    def test_render_contract_preserves_prefix_and_removes_template_body(self):
        doc = Document()
        title = doc.add_paragraph("模板标题")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("摘  要：[填写报告摘要。]")
        doc.add_paragraph("报告单位：示例单位")
        doc.add_paragraph("一、模板章节")
        doc.add_paragraph("模板示例正文。")
        path = self._save(doc)
        schema = compile_template(path)
        rendered = Document(path)

        title_anchor = _prepare_render_body(rendered, schema)

        self.assertIsNotNone(title_anchor)
        self.assertEqual([item.text for item in rendered.paragraphs], ["模板标题", "报告单位：示例单位"])

    def test_numbering_strategy_reads_ooxml_role_levels(self):
        schema = {
            "style": {
                "roles": {},
                "numbering": {
                    "role_levels": {
                        "heading_1": {"number_format": "decimal", "level_text": "%1 "},
                        "heading_2": {"number_format": "decimal", "level_text": "%1.%2 "},
                    }
                },
            }
        }
        strategy = detect_numbering_strategy(schema)
        self.assertEqual(format_heading(1, [2], "研究方法", strategy), "2 研究方法")
        self.assertEqual(format_heading(2, [2, 3], "实验设计", strategy), "2.3 实验设计")

    def test_runtime_style_retrieval_refuses_irrelevant_fillers(self):
        bank = [{
            "content": "概述研究范围和问题边界。",
            "purpose": "介绍背景",
            "sample_type": "opening",
            "source_report": "A",
            "section": "背景",
            "approved": True,
        }]
        selected = select_exemplars(
            bank,
            {"purpose": "形成风险判断", "sample_type": "risk", "keywords": ["供应链"]},
            limit=3,
        )
        self.assertEqual(selected, [])

    def test_style_metrics_expose_material_realization_signals(self):
        metrics = aggregate_style_metrics([{
            "content": "根据《专项规划》，项目于2026年完成三项验证，这表明实施路径已具备基础条件。",
            "structural_role": "body",
            "source_report": "A",
        }])
        realization = metrics["material_realization"]
        self.assertEqual(realization["concrete_detail_ratio"], 1.0)
        self.assertEqual(realization["explicit_attribution_ratio"], 1.0)
        self.assertEqual(realization["fact_judgment_combination_ratio"], 1.0)


if __name__ == "__main__":
    unittest.main()
